import docx
from markdownify import markdownify as md
import os

def read_docx(file_path):
    doc = docx.Document(file_path)
    elements = []

    for para in doc.paragraphs:
        elements.append(para.text)

    for table in doc.tables:
        for row in table.rows:
            row_data = []
            for cell in row.cells:
                row_data.append(cell.text)
            elements.append('| ' + ' | '.join(row_data) + ' |')
        
        # Adding a separator line for markdown tables
        if len(table.rows) > 0:
            header = table.rows[0]
            separator = '| ' + ' | '.join(['---']*len(header.cells)) + ' |'
            elements.append(separator)
    
    return '\n'.join(elements)

def convert_to_markdown(docx_content):
    return md(docx_content, heading_style="ATX")

def save_markdown(markdown_content, output_path):
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(markdown_content)

def convert_docx_to_markdown(input_file, output_file):
    docx_content = read_docx(input_file)
    markdown_content = convert_to_markdown(docx_content)
    save_markdown(markdown_content, output_file)

# Exemplo de uso
input_file = 'documentacao.docx'
output_file = 'documentacao.md'

# Verifique se o arquivo de entrada existe
if not os.path.isfile(input_file):
    print(f"Arquivo de entrada {input_file} não encontrado.")
else:
    convert_docx_to_markdown(input_file, output_file)
    print(f"Arquivo convertido e salvo em {output_file}")